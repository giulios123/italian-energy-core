from __future__ import annotations

import io
from datetime import UTC, datetime
from decimal import Decimal
from hashlib import sha256
from typing import Any, ClassVar
from urllib.request import Request

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from italian_energy.arera import (
    ARERA_DOMESTIC_2026_URL,
    AreraChargeRole,
    AreraChargeValue,
    AreraCustomerSegment,
    AreraDomesticElectricityImporter,
    AreraFetchError,
    AreraHeader,
    AreraHttpResponse,
    AreraImportError,
    AreraRegulatoryBundle,
    AreraSourceLocator,
    RawAreraSnapshot,
)
from italian_energy.arera import importer as importer_module
from italian_energy.domain.money import RateUnit, UnitRate
from italian_energy.domain.regulatory import BillingCategory, BillingQuota, VerificationStatus
from italian_energy.domain.time import DatePeriod

RETRIEVED_AT = datetime(2026, 9, 7, 9, 0, tzinfo=UTC)
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _set_headers(ws: Worksheet, marker_row: int) -> None:
    ws[f"C{marker_row + 1}"] = "Vendita di energia elettrica"
    ws[f"I{marker_row + 1}"] = "Tariffa per l'uso della rete elettrica"
    ws[f"L{marker_row + 1}"] = "Oneri generali di sistema"
    for column, value in {
        "C": "dispacciamento",
        "D": "\u03c31",
        "E": "\u03c32",
        "F": "\u03c33",
        "G": "UC3",
        "H": "UC6",
        "J": "ASOS",
        "K": "ARIM",
    }.items():
        ws[f"{column}{marker_row + 2}"] = value


def _set_segment(
    ws: Worksheet,
    marker_row: int,
    resident: bool,
    mismatch: bool,
    *,
    bad_network_total: bool = False,
    bad_group: bool = False,
    bad_component: bool = False,
    bad_quota: bool = False,
    non_numeric: bool = False,
    bad_format: bool = False,
    negative: bool = False,
    duplicate_segment: bool = False,
    extra_component: bool = False,
) -> None:
    ws[f"B{marker_row}"] = (
        "ABITAZIONI DI RESIDENZE ANAGRAFICHE"
        if resident
        else "ABITAZIONI DIVERSE DA RESIDENZE ANAGRAFICHE"
    )
    _set_headers(ws, marker_row)
    if extra_component:
        ws[f"M{marker_row + 2}"] = "NUOVA"
    if bad_group:
        ws[f"C{marker_row + 1}"] = "schema cambiato"
    if bad_component:
        ws[f"D{marker_row + 2}"] = "sigma1"
    if duplicate_segment and not resident:
        ws[f"B{marker_row}"] = "ABITAZIONI DI RESIDENZE ANAGRAFICHE"
    row_values: dict[int, dict[str, Any]] = {
        marker_row + 4: {
            "B": "Quota energia (euro/kWh)",
            "C": "invalid" if non_numeric else (-0.019902 if negative else 0.019902),
            "D": "-",
            "E": "-",
            "F": 0.0119,
            "G": 0.00276,
            "H": 0.00007,
            "I": 0.014731 if bad_network_total else 0.01473,
            "J": 0.028657,
            "K": 0.001638,
            "L": 0.030296 if mismatch else 0.030295,
        },
        marker_row + 5: {
            "B": "quota sconosciuta" if bad_quota else "Quota fissa (euro/anno)",
            "C": "-",
            "D": 23.04,
            "E": "-",
            "F": "-",
            "G": "-",
            "H": "-",
            "I": 23.04,
            "J": 88.752 if not resident else "-",
            "K": 0 if not resident else "-",
            "L": 88.752 if not resident else 0,
        },
        marker_row + 6: {
            "B": "Quota potenza (euro/kW/anno)",
            "C": "-",
            "D": "-",
            "E": 23.52,
            "F": "-",
            "G": "-",
            "H": 0.1988,
            "I": 23.7188,
            "J": "-",
            "K": "-",
            "L": 0,
        },
    }
    if mismatch:
        row_values[marker_row + 4]["L"] = 0.030296
    for row, values in row_values.items():
        for column, value in values.items():
            cell = ws[f"{column}{row}"]
            cell.value = value
            if not isinstance(value, (int, float)):
                continue
            if row == marker_row + 4:
                cell.number_format = "#,##0.000000"
            elif row == marker_row + 5:
                cell.number_format = "#,##0.00"
            elif row == marker_row + 6:
                cell.number_format = "#,##0.0000"
    if bad_format:
        ws[f"C{marker_row + 4}"].number_format = "General"


def _workbook_bytes(
    months: tuple[str, ...] = ("giugno 2026", "maggio 2026"),
    *,
    mismatch: bool = False,
    formula: bool = False,
    extra_sheet: bool = False,
    bad_network_total: bool = False,
    bad_group: bool = False,
    bad_component: bool = False,
    bad_quota: bool = False,
    non_numeric: bool = False,
    bad_format: bool = False,
    negative: bool = False,
    duplicate_segment: bool = False,
    extra_component: bool = False,
) -> bytes:
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)
    for month in months:
        worksheet = workbook.create_sheet(month)
        worksheet["B2"] = month.title()
        worksheet["B4"] = "CLIENTI DOMESTICI CON FORNITURA NEL MERCATO LIBERO"
        _set_segment(
            worksheet,
            10,
            True,
            mismatch,
            bad_network_total=bad_network_total,
            bad_group=bad_group,
            bad_component=bad_component,
            bad_quota=bad_quota,
            non_numeric=non_numeric,
            bad_format=bad_format,
            negative=negative,
            duplicate_segment=duplicate_segment,
            extra_component=extra_component,
        )
        _set_segment(
            worksheet,
            20,
            False,
            mismatch,
            bad_network_total=bad_network_total,
            bad_group=bad_group,
            bad_component=bad_component,
            bad_quota=bad_quota,
            non_numeric=non_numeric,
            bad_format=bad_format,
            negative=negative,
            duplicate_segment=duplicate_segment,
            extra_component=extra_component,
        )
        if formula:
            worksheet["C14"] = "=1+1"
    if extra_sheet:
        workbook.create_sheet("nota 2026")["A1"] = "schema cambiato"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _response(
    content: bytes,
    *,
    status: int = 200,
    url: str = ARERA_DOMESTIC_2026_URL,
    content_type: str = MIME,
) -> AreraHttpResponse:
    return AreraHttpResponse(
        status=status,
        url=url,
        headers=(AreraHeader(name="Content-Type", value=content_type),),
        content=content,
    )


class _TypedFakeTransport:
    def __init__(self, response: AreraHttpResponse):
        self.response = response

    def get(self, url: str, timeout_seconds: float) -> AreraHttpResponse:
        return self.response


class _ErrorTransport:
    def get(self, url: str, timeout_seconds: float) -> AreraHttpResponse:
        raise TimeoutError("synthetic timeout")


def test_parse_bytes_returns_unverified_source_faithful_bundle() -> None:
    content = _workbook_bytes()
    result = AreraDomesticElectricityImporter().parse_bytes(content, RETRIEVED_AT)

    assert result.status == VerificationStatus.UNVERIFIED
    assert result.bundle is not None
    assert result.bundle.status == VerificationStatus.UNVERIFIED
    assert result.snapshot.content == content
    assert result.snapshot.snapshot_id == f"arera-snapshot:{sha256(content).hexdigest()}"
    assert result.bundle.bundle_id == f"arera-bundle:{sha256(content).hexdigest()}"
    assert len(result.bundle.charges) == 64
    assert {charge.segment.value for charge in result.bundle.charges} == {
        "residenza_anagrafica",
        "diversa_da_residenza_anagrafica",
    }
    assert all(isinstance(charge.rate.amount, Decimal) for charge in result.bundle.charges)
    assert {charge.role.value for charge in result.bundle.charges} == {"atomic", "total"}
    assert result.bundle.model_dump_json()


def test_official_fetch_promotes_values_to_verified() -> None:
    content = _workbook_bytes()
    importer = AreraDomesticElectricityImporter(
        transport=_TypedFakeTransport(_response(content)),
        clock=lambda: RETRIEVED_AT,
    )

    result = importer.fetch_and_import()

    assert result.status == VerificationStatus.VERIFIED
    assert result.bundle is not None
    assert all(charge.status == VerificationStatus.VERIFIED for charge in result.bundle.charges)
    assert result.snapshot.content_type == MIME
    assert result.bundle.provenance[0].sha256 == sha256(content).hexdigest()


def test_source_values_and_totals_are_preserved_without_billing_mapping() -> None:
    result = AreraDomesticElectricityImporter().parse_bytes(_workbook_bytes(), RETRIEVED_AT)
    assert result.bundle is not None
    charges = {charge.code: charge for charge in result.bundle.charges}

    energy = charges["arera:2026:05:residenza_anagrafica:network_total:consumption"]
    dispatch = charges["arera:2026:05:residenza_anagrafica:CDISPD:consumption"]
    assert energy.rate.amount == Decimal("0.014730")
    assert dispatch.rate.amount == Decimal("0.019902")
    assert energy.role.value == "total"
    assert dispatch.role.value == "atomic"
    assert energy.provenance[0].effective_period is not None
    assert "arera:2026:05:residenza_anagrafica:sigma2:fixed" not in charges
    assert charges[
        "arera:2026:05:diversa_da_residenza_anagrafica:ARIM:fixed"
    ].rate.amount == Decimal("0.00")


@pytest.mark.parametrize(
    ("kwargs", "code"),
    [
        ({"months": ("gennaio 2026", "marzo 2026")}, "MONTH_GAP"),
        ({"extra_sheet": True}, "UNKNOWN_SHEET"),
        ({"extra_component": True}, "HEADER_COMPONENT"),
        ({"mismatch": True}, "SYSTEM_TOTAL"),
        ({"formula": True}, "FORMULA"),
    ],
)
def test_schema_drift_returns_review_required_without_bundle(
    kwargs: dict[str, Any], code: str
) -> None:
    result = AreraDomesticElectricityImporter().parse_bytes(_workbook_bytes(**kwargs), RETRIEVED_AT)

    assert result.status == VerificationStatus.REVIEW_REQUIRED
    assert result.bundle is None
    assert code in {diagnostic.code for diagnostic in result.diagnostics}
    assert result.snapshot.content


@pytest.mark.parametrize(
    ("kwargs", "code"),
    [
        ({"bad_network_total": True}, "NETWORK_TOTAL"),
        ({"bad_group": True}, "HEADER_GROUP"),
        ({"bad_component": True}, "HEADER_COMPONENT"),
        ({"bad_quota": True}, "QUOTA_ROWS"),
        ({"non_numeric": True}, "NON_NUMERIC"),
        ({"bad_format": True}, "NUMBER_FORMAT"),
        ({"negative": True}, "ECONOMIC_VALUE"),
        ({"duplicate_segment": True}, "DUPLICATE_SEGMENT"),
    ],
)
def test_value_and_layout_diagnostics_are_fail_closed(kwargs: dict[str, Any], code: str) -> None:
    result = AreraDomesticElectricityImporter().parse_bytes(_workbook_bytes(**kwargs), RETRIEVED_AT)
    assert result.status == VerificationStatus.REVIEW_REQUIRED
    assert result.bundle is None
    assert code in {diagnostic.code for diagnostic in result.diagnostics}


@pytest.mark.parametrize(
    ("status", "url", "content_type"),
    [
        (404, ARERA_DOMESTIC_2026_URL, MIME),
        (200, "https://evil.example/redirect.xlsx", MIME),
        (200, ARERA_DOMESTIC_2026_URL, "text/html"),
    ],
)
def test_fetch_rejects_untrusted_transport_response(
    status: int, url: str, content_type: str
) -> None:
    importer = AreraDomesticElectricityImporter(
        transport=_TypedFakeTransport(
            _response(_workbook_bytes(), status=status, url=url, content_type=content_type)
        )
    )
    with pytest.raises(AreraFetchError):
        importer.fetch_and_import()


def test_input_limits_and_snapshot_validation_are_fail_closed() -> None:
    with pytest.raises(AreraImportError):
        AreraDomesticElectricityImporter().parse_bytes(b"x" * 5_000_001, RETRIEVED_AT)

    with pytest.raises(ValueError, match="sha256"):
        RawAreraSnapshot(
            content=b"raw",
            retrieved_at=RETRIEVED_AT,
            content_type="application/octet-stream",
            sha256="0" * 64,
        )

    with pytest.raises(ValueError, match="timezone"):
        RawAreraSnapshot(
            content=b"raw",
            retrieved_at=datetime(2026, 9, 7, 9, 0),
            content_type="application/octet-stream",
            sha256=sha256(b"raw").hexdigest(),
        )


def test_only_2026_is_supported_and_timeout_must_be_positive() -> None:
    importer = AreraDomesticElectricityImporter(transport=_TypedFakeTransport(_response(b"")))
    with pytest.raises(Exception, match="2026"):
        importer.fetch_and_import(year=2025)
    with pytest.raises(ValueError, match="timeout"):
        importer.fetch_and_import(timeout_seconds=0)
    with pytest.raises(AreraFetchError, match="synthetic timeout"):
        AreraDomesticElectricityImporter(transport=_ErrorTransport()).fetch_and_import()


def test_malformed_zip_and_empty_workbook_are_review_required() -> None:
    malformed = AreraDomesticElectricityImporter().parse_bytes(b"not an xlsx", RETRIEVED_AT)
    assert malformed.status == VerificationStatus.REVIEW_REQUIRED
    assert malformed.diagnostics[0].code == "IMPORT_ERROR"

    empty_bytes = AreraDomesticElectricityImporter().parse_bytes(b"", RETRIEVED_AT)
    assert empty_bytes.status == VerificationStatus.REVIEW_REQUIRED
    assert empty_bytes.diagnostics[0].code == "IMPORT_ERROR"

    empty = AreraDomesticElectricityImporter().parse_bytes(
        _workbook_bytes(months=("nota 2026",)), RETRIEVED_AT
    )
    assert empty.status == VerificationStatus.REVIEW_REQUIRED
    assert {item.code for item in empty.diagnostics} >= {"UNKNOWN_SHEET", "NO_MONTHS"}


def test_duplicate_months_and_invalid_timeout_variants() -> None:
    content = _workbook_bytes(months=("maggio 2026",))
    parsed = AreraDomesticElectricityImporter().parse_bytes(content, RETRIEVED_AT)
    raw_sheets, workbook = importer_module._read_workbook(content)
    diagnostics: list[Any] = []
    importer_module.AreraDomesticElectricityImporter()._parse_sheets(
        parsed.snapshot,
        (raw_sheets[0], raw_sheets[0]),
        workbook,
        diagnostics,
    )
    assert "DUPLICATE_MONTH" in {item.code for item in diagnostics}

    importer = AreraDomesticElectricityImporter(
        transport=_TypedFakeTransport(_response(_workbook_bytes()))
    )
    with pytest.raises(ValueError, match="timeout"):
        importer.fetch_and_import(timeout_seconds=float("inf"))


def test_archive_guards_and_decimal_helpers() -> None:
    import zipfile

    def archive_with(names: tuple[str, ...]) -> bytes:
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            for name in names:
                archive.writestr(name, b"x")
        return output.getvalue()

    with pytest.raises(AreraImportError, match="required"):
        importer_module._validate_archive(archive_with(("other.txt",)))
    with pytest.raises(AreraImportError, match="macros"):
        importer_module._validate_archive(
            archive_with(
                ("[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "xl/vbaProject.bin")
            )
        )
    with pytest.raises(AreraImportError, match="unsafe"):
        importer_module._validate_archive(
            archive_with(("[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "../escape"))
        )
    with pytest.raises(AreraImportError, match="duplicate"):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("[Content_Types].xml", b"x")
            archive.writestr("[Content_Types].xml", b"y")
        importer_module._validate_archive(output.getvalue())

    assert importer_module._decimal_places("#,##0.000000") == 6
    assert importer_module._decimal_places("#,##0.00") == 2
    assert importer_module._decimal_places("General") is None


def test_economic_cell_parser_reports_all_boundary_errors() -> None:
    diagnostics: list[Any] = []
    raw = importer_module._RawCell(ref="A1", kind="s", value="abc", formula=False)
    assert importer_module._economic_cell_value(raw, "General", diagnostics, "x", "A1") is None
    raw = importer_module._RawCell(ref="A1", kind=None, value=None, formula=False)
    assert importer_module._economic_cell_value(raw, "General", diagnostics, "x", "A1") is None
    raw = importer_module._RawCell(ref="A1", kind=None, value="not-decimal", formula=False)
    assert importer_module._economic_cell_value(raw, "#,##0.00", diagnostics, "x", "A1") is None
    raw = importer_module._RawCell(ref="A1", kind=None, value="Infinity", formula=False)
    assert importer_module._economic_cell_value(raw, "#,##0.00", diagnostics, "x", "A1") is None
    raw = importer_module._RawCell(ref="A1", kind=None, value="1.234", formula=True)
    assert importer_module._economic_cell_value(raw, "#,##0.00", diagnostics, "x", "A1") is None
    assert {item.code for item in diagnostics} >= {
        "NON_NUMERIC",
        "MISSING_VALUE",
        "DECIMAL",
        "ECONOMIC_VALUE",
        "FORMULA",
    }
    precision_diagnostics: list[Any] = []
    precise = importer_module._RawCell(ref="A1", kind=None, value="1.23456789", formula=False)
    assert importer_module._economic_cell_value(
        precise, "#,##0.00000", precision_diagnostics, "x", "A1"
    ) == Decimal("1.23457")
    assert importer_module._quantize_for_format(Decimal("23.7188"), "#,##0.00") == Decimal("23.72")


def test_total_validation_rounds_sum_to_declared_total_format() -> None:
    diagnostics: list[Any] = []
    importer_module.AreraDomesticElectricityImporter._validate_totals(
        {
            BillingQuota.POWER: {
                "E": Decimal("23.52"),
                "H": Decimal("0.198800"),
                "I": Decimal("23.72"),
                "L": Decimal("0.00"),
            }
        },
        {
            BillingQuota.POWER: {
                "E": "#,##0.00",
                "H": "#,##0.000000",
                "I": "#,##0.00",
                "L": "#,##0.00",
            }
        },
        BillingQuota.POWER,
        "gennaio 2026",
        16,
        diagnostics,
    )
    assert diagnostics == []


def test_bundle_models_reject_inconsistent_values() -> None:
    period = DatePeriod(start=datetime(2026, 5, 1).date(), end=datetime(2026, 6, 1).date())
    locator = AreraSourceLocator(sheet="maggio 2026", cell="C14", label="dispacciamento")
    charge = AreraChargeValue(
        code="one",
        segment=AreraCustomerSegment.RESIDENT,
        component_code="CDISPD",
        category=BillingCategory.SALES,
        quota=BillingQuota.CONSUMPTION,
        rate=UnitRate(amount=Decimal("0.1"), unit=RateUnit.EUR_PER_KWH),
        validity=period,
        role=AreraChargeRole.ATOMIC,
        source=locator,
        status=VerificationStatus.UNVERIFIED,
    )
    with pytest.raises(ValueError, match="requires charges"):
        AreraRegulatoryBundle(
            bundle_id="empty",
            schema_version="v1",
            dataset_version="2026",
            validity=period,
            charges=(),
            status=VerificationStatus.UNVERIFIED,
        )
    with pytest.raises(ValueError, match="unique"):
        AreraRegulatoryBundle(
            bundle_id="duplicate",
            schema_version="v1",
            dataset_version="2026",
            validity=period,
            charges=(charge, charge),
            status=VerificationStatus.UNVERIFIED,
        )


class _BufferedResponse:
    status = 200
    headers: ClassVar[dict[str, str]] = {"Content-Type": MIME, "ETag": "test"}

    def __init__(self, content: bytes, url: str = ARERA_DOMESTIC_2026_URL):
        self.content = content
        self.position = 0
        self.url = url

    def __enter__(self) -> _BufferedResponse:
        return self

    def __exit__(self, *_args: object) -> None:
        return None

    def read(self, size: int) -> bytes:
        result = self.content[self.position : self.position + size]
        self.position += len(result)
        return result

    def geturl(self) -> str:
        return self.url


class _BufferedOpener:
    def __init__(self, response: _BufferedResponse | Exception):
        self.response = response

    def open(self, _request: object, timeout: float) -> _BufferedResponse:
        assert timeout == 1.0
        if isinstance(self.response, Exception):
            raise self.response
        return self.response


def test_urllib_transport_and_redirect_guard(monkeypatch: pytest.MonkeyPatch) -> None:
    response = _BufferedResponse(b"abc")
    monkeypatch.setattr(importer_module, "build_opener", lambda _handler: _BufferedOpener(response))
    fetched = importer_module._UrllibTransport().get(ARERA_DOMESTIC_2026_URL, 1.0)
    assert fetched.content == b"abc"
    assert fetched.status == 200

    guard = importer_module._RedirectGuard()
    redirected = guard.redirect_request(
        Request(ARERA_DOMESTIC_2026_URL),
        None,
        302,
        "found",
        {},
        ARERA_DOMESTIC_2026_URL,
    )
    assert redirected is not None
    with pytest.raises(AreraFetchError, match="allowlisted"):
        guard.redirect_request(
            Request(ARERA_DOMESTIC_2026_URL),
            None,
            302,
            "found",
            {},
            "https://evil.example/xlsx",
        )


def test_urllib_transport_wraps_network_errors_and_read_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        importer_module,
        "build_opener",
        lambda _handler: _BufferedOpener(RuntimeError("offline")),
    )
    with pytest.raises(AreraFetchError, match="fetch failed"):
        importer_module._UrllibTransport().get(ARERA_DOMESTIC_2026_URL, 1.0)

    class Reader:
        def read(self, size: int) -> bytes:
            return b"x" * size

    with pytest.raises(AreraFetchError, match="maximum content"):
        importer_module._read_limited(Reader(), 2)


def test_shared_strings_and_sheet_parser_paths() -> None:
    import zipfile

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            "<si><t>hello</t></si></sst>",
        )
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        assert importer_module._read_shared_strings(archive) == ("hello",)

    sheet = importer_module._read_sheet(
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c>'
        b'<c r="B1" t="inlineStr"><is><t>world</t></is></c></row></sheetData></worksheet>',
        "test",
        ("hello",),
    )
    assert sheet.cells["A1"].value == "hello"
    assert sheet.cells["B1"].value == "world"
    with pytest.raises(AreraImportError, match="shared string"):
        importer_module._read_sheet(
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<sheetData><row r="1"><c r="A1" t="s"><v>4</v></c></row></sheetData></worksheet>',
            "test",
            ("hello",),
        )
