from __future__ import annotations

import io
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from italian_energy.arera.discovery import RegulatoryRegistryChannel
from italian_energy.arera.importer import SCHEMA_VERSION as WORKBOOK_LAYOUT
from italian_energy.arera.models import AreraCustomerSegment
from italian_energy.arera.rollover_models import RegulatoryFactFamily, RegulatoryRolloverReason
from italian_energy.arera.rollover_parsers import (
    ADM_EXCISE_DOMESTIC_KIND,
    ADM_EXCISE_DOMESTIC_PARSER_ID,
    ADM_EXCISE_DOMESTIC_PARSER_VERSION,
    ADM_EXCISE_LAYOUT_VERSION,
    ARERA_343_CONFIRMATION_DOCUMENT_KIND,
    ARERA_343_CONFIRMATION_LAYOUT_VERSION,
    ARERA_343_CONFIRMATION_PARSER_ID,
    ARERA_343_CONFIRMATION_PARSER_VERSION,
    ARERA_343_PDF_URL,
    ARERA_DOMESTIC_WORKBOOK_KIND,
    ARERA_DOMESTIC_WORKBOOK_MIME,
    ARERA_DOMESTIC_WORKBOOK_PARSER_ID,
    ARERA_DOMESTIC_WORKBOOK_PARSER_VERSION,
    NORMATTIVA_VAT_ART16_KIND,
    NORMATTIVA_VAT_LAYOUT_VERSION,
    NORMATTIVA_VAT_RATE_PARSER_ID,
    NORMATTIVA_VAT_TABLE_A_KIND,
    NORMATTIVA_VAT_TABLE_PARSER_ID,
    AreraDomesticWorkbookParser,
    ParserDisposition,
    RegulatoryParseResult,
    RegulatoryParserRegistry,
    RegulatorySourceDocument,
    default_regulatory_parser_registry,
)
from italian_energy.arera.rollover_sources import AcquiredOfficialBytes
from italian_energy.domain.money import RateUnit
from italian_energy.domain.regulatory import BillingQuota, VerificationStatus
from italian_energy.domain.time import DatePeriod

FETCHED_AT = datetime(2026, 9, 30, 12, tzinfo=UTC)
PUBLISHED_AT = date(2026, 9, 25)
SOURCE_URL = (
    "https://www.arera.it/fileadmin/area_operatori/prezzi_e_tariffe/"
    "Corrispettivi_libero_elettrico_domestico_2026.xlsx"
)
VAT_TABLE_FIXTURE = (
    Path(__file__).parents[1]
    / "fixtures/regulatory_rollover/normattiva-dpr633-table-a-synthetic.html"
)
VAT_ART16_FIXTURE = (
    Path(__file__).parents[1]
    / "fixtures/regulatory_rollover/normattiva-dpr633-art16-synthetic.html"
)


def _synthetic_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "ottobre 2026"
    sheet["B2"] = "Ottobre 2026"
    sheet["B4"] = "CLIENTI DOMESTICI CON FORNITURA NEL MERCATO LIBERO"
    for marker_row, resident in ((10, True), (20, False)):
        sheet[f"B{marker_row}"] = (
            "ABITAZIONI DI RESIDENZE ANAGRAFICHE"
            if resident
            else "ABITAZIONI DIVERSE DA RESIDENZE ANAGRAFICHE"
        )
        for column, title in {
            "C": "Vendita di energia elettrica",
            "I": "Tariffa per l'uso della rete elettrica",
            "L": "Oneri generali di sistema",
        }.items():
            sheet[f"{column}{marker_row + 1}"] = title
        for column, title in {
            "C": "dispacciamento",
            "D": "\u03c31",
            "E": "\u03c32",
            "F": "\u03c33",
            "G": "UC3",
            "H": "UC6",
            "J": "ASOS",
            "K": "ARIM",
        }.items():
            sheet[f"{column}{marker_row + 2}"] = title

        rows: tuple[tuple[str, dict[str, Any]], ...] = (
            (
                "Quota energia (euro/kWh)",
                {
                    "C": 0.019902,
                    "D": "-",
                    "E": "-",
                    "F": 0.0119,
                    "G": 0.00276,
                    "H": 0.00007,
                    "I": 0.01473,
                    "J": 0.028657,
                    "K": 0.001638,
                    "L": 0.030295,
                },
            ),
            (
                "Quota fissa (euro/anno)",
                {
                    "C": "-",
                    "D": 23.04,
                    "E": "-",
                    "F": "-",
                    "G": "-",
                    "H": "-",
                    "I": 23.04,
                    "J": "-" if resident else 88.752,
                    "K": "-" if resident else 0,
                    "L": 0 if resident else 88.752,
                },
            ),
            (
                "Quota potenza (euro/kW/anno)",
                {
                    "C": "-",
                    "D": "-",
                    "E": 23.52,
                    "F": "-",
                    "G": "-",
                    "H": 0.1988,
                    "I": 23.7188,
                    "J": "-",
                    "K": "-",
                    "L": 0,
                },
            ),
        )
        for offset, (quota_label, values) in enumerate(rows, start=4):
            row = marker_row + offset
            sheet[f"B{row}"] = quota_label
            for column in "CDEFGHIJKL":
                cell = sheet[f"{column}{row}"]
                cell.value = values[column]
                cell.number_format = (
                    "#,##0.000000" if offset == 4 else "#,##0.00" if offset == 5 else "#,##0.0000"
                )
    for month_name in ("novembre 2026", "dicembre 2026"):
        copied = workbook.copy_worksheet(sheet)
        copied.title = month_name
        copied["B2"] = month_name.title()
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _source(
    *,
    body: bytes | None = None,
    layout_version: str = WORKBOOK_LAYOUT,
    document_kind: str = ARERA_DOMESTIC_WORKBOOK_KIND,
) -> RegulatorySourceDocument:
    content = _synthetic_workbook() if body is None else body
    acquired = AcquiredOfficialBytes(
        url=SOURCE_URL,
        body=content,
        sha256=sha256(content).hexdigest(),
        fetched_at=FETCHED_AT,
    )
    return RegulatorySourceDocument(
        acquired=acquired,
        source_id="arera-workbook-2026-q4-fixture",
        channel=RegulatoryRegistryChannel.ARERA_TARIFFS,
        act_id="R-fixture-2026-01",
        document_id="synthetic workbook fixture",
        published_at=PUBLISHED_AT,
        content_type=ARERA_DOMESTIC_WORKBOOK_MIME,
        document_kind=document_kind,
        layout_version=layout_version,
    )


def _normattiva_vat_source(
    *,
    table: bool,
    body: bytes | None = None,
    layout_version: str = NORMATTIVA_VAT_LAYOUT_VERSION,
    source_id: str | None = None,
    url_date: date = date(2026, 10, 5),
) -> RegulatorySourceDocument:
    if table:
        source_id = source_id or "vat_dpr_633"
        urn = "urn:nir:stato:decreto.presidente.della.repubblica:1972-10-26;633:1~art1"
        document_id = "DPR-633-1972-Tabella-A-Parte-III-n-103"
        kind = NORMATTIVA_VAT_TABLE_A_KIND
        page_bytes = VAT_TABLE_FIXTURE.read_bytes()
    else:
        source_id = source_id or "vat_dpr_633_art16"
        urn = "urn:nir:stato:decreto.presidente.della.repubblica:1972-10-26;633~art16"
        document_id = "DPR-633-1972-art-16-aliquote"
        kind = NORMATTIVA_VAT_ART16_KIND
        page_bytes = VAT_ART16_FIXTURE.read_bytes()
    content = page_bytes if body is None else body
    url = f"https://www.normattiva.it/uri-res/N2Ls?{urn}!vig={url_date.isoformat()}"
    return RegulatorySourceDocument(
        acquired=AcquiredOfficialBytes(
            url=url,
            body=content,
            sha256=sha256(content).hexdigest(),
            fetched_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
            content_type="text/html; charset=UTF-8",
        ),
        source_id=source_id,
        channel=RegulatoryRegistryChannel.NORMATTIVA_UPDATES,
        act_id="DPR 633/1972",
        document_id=document_id,
        published_at=date(1972, 10, 26),
        content_type="text/html; charset=UTF-8",
        document_kind=kind,
        layout_version=layout_version,
    )


def _synthetic_adm_pdf(
    *lines: str,
    page_count: int = 11,
    positions_override: dict[str, tuple[int, int]] | None = None,
) -> bytes:
    """Build a tiny text PDF for offline layout tests, without network fixtures."""

    positions = {
        "Aggiornamento al 18 settembre 2026": (316, 537),
        "ACCISE SULL' ENERGIA ELETTRICA": (284, 502),
        "per qualsiasi applicazione nelle abitazioni": (38, 474),
        "nelle abitazioni": (62, 460),
        "€ 0,0227 per ogni kWh": (175, 475),
        "€ 0,0228 per ogni kWh": (175, 475),
        "D.M. 30/12/2011": (202, 458),
        "Pagina 5 di 11": (388, 19),
    }
    commands: list[str] = []
    overrides = positions_override or {}
    for line in lines:
        if line == "per qualsiasi applicazione nelle abitazioni":
            current_positions: tuple[tuple[int, int, str], ...] = (
                (
                    *overrides.get("per qualsiasi applicazione", (38, 474)),
                    "per qualsiasi applicazione",
                ),
                (*overrides.get("nelle abitazioni", (62, 460)), "nelle abitazioni"),
            )
        else:
            if line.startswith("Aggiornamento al "):
                default_position = (316, 537)
            elif line.startswith("€ 0,") and line.endswith(" per ogni kWh"):
                default_position = (175, 475)
            else:
                default_position = positions.get(line, (50, 740 - len(commands) * 18))
            x, y = overrides.get(line, default_position)
            current_positions = ((x, y, line),)
        for x, y, text in current_positions:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.extend((f"BT /F1 11 Tf {x} {y} Td", f"({escaped}) Tj ET"))
    stream = "\n".join(commands).encode("cp1252")
    font_id = 3 + page_count
    content_start = font_id + 1
    kids = " ".join(f"{page_id} 0 R" for page_id in range(3, font_id))
    objects_list = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode(),
    ]
    for page_offset in range(page_count):
        content_id = content_start + page_offset
        objects_list.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {content_id} 0 R >>".encode()
        )
    objects_list.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    for page_offset in range(page_count):
        page_stream = stream if page_offset == 4 else b""
        objects_list.append(
            b"<< /Length "
            + str(len(page_stream)).encode()
            + b" >>\nstream\n"
            + page_stream
            + b"\nendstream"
        )
    objects = tuple(objects_list)
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)


def _synthetic_343_pdf(
    *,
    article_text: str | None = None,
    article_replace: tuple[str, str] | None = None,
    cover_text: str | None = None,
    final_page_text: str | None = None,
    page_count: int = 8,
    empty_page: int | None = None,
) -> bytes:
    """Build an extractable eight-page ARERA confirmation act for offline tests."""

    cover = cover_text or (
        "DELIBERAZIONE 29 SETTEMBRE 2026 343/2026/R/COM AGGIORNAMENTO DAL 1 OTTOBRE 2026 "
        "DELLE COMPONENTI TARIFFARIE"
    )
    article = article_text or (
        "Articolo 1 Componenti tariffarie relative al settore elettrico "
        "1.1 I valori della componente tariffaria ASOS in vigore a decorrere dal 1 luglio "
        "2026 per le utenze che non sono nella titolarita di imprese a forte consumo di "
        "energia elettrica di cui alla Tabella 1 allegata alla deliberazione 227/2026/R/com "
        "sono confermati. "
        "1.2 I valori della componente tariffaria ASOS in vigore a decorrere dal 1 luglio "
        "2026 per le utenze che sono nella titolarita di imprese a forte consumo di energia "
        "elettrica di cui alle Tabelle 2, 3, 4 e 5 allegate alla deliberazione 227/2026/R/com "
        "sono confermati. "
        "1.3 I valori della componente tariffaria ARIM in vigore a decorrere dal 1 gennaio "
        "2026 di cui alla Tabella 6 allegata alla deliberazione 588/2025/R/com sono confermati. "
        "1.4 I valori delle componenti tariffarie UC3 e UC6 in vigore a decorrere dal 1 gennaio "
        "2026 di cui alla Tabella 7 allegata alla deliberazione 588/2025/R/com sono confermati. "
        "1.5 I valori delle componenti tariffarie ASOS ARIM UC3 e UC6 in vigore a decorrere "
        "dal 1 luglio 2026 per il soggetto di cui al comma 36.1 del TIPPI di cui alla Tabella "
        "6 allegata alla deliberazione 227/2026/R/com sono confermati. "
        "1.6 Le percentuali di ripartizione della componente ARIM di cui al comma 3.7 del TIPPI "
        "in vigore a decorrere dal 1 gennaio 2026 di cui al comma 1.6 della deliberazione "
        "588/2025/R/com sono confermati. "
        "1.7 Il 100% della componente ASOS e da destinare al Conto per nuovi impianti da fonti "
        "rinnovabili e assimilate di cui al comma 10.1 lettera b del TIPPI."
    )
    if article_replace is not None:
        article = article.replace(*article_replace)
    final_page = final_page_text or (
        "4.2 Il presente provvedimento e pubblicato sul sito internet dell Autorita "
        "ed entra in vigore dal 1 ottobre 2026."
    )
    page_texts = [
        f"Premesse della deliberazione 343/2026/R/com, pagina {index + 1}."
        for index in range(page_count)
    ]
    if page_count:
        page_texts[0] = cover
    if page_count > 5:
        page_texts[5] = article
    if page_count > 7:
        page_texts[7] = final_page
    if empty_page is not None and empty_page < page_count:
        page_texts[empty_page] = ""

    def escaped(value: str) -> bytes:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("cp1252")

    page_ids = tuple(range(3, 3 + page_count))
    font_id = 3 + page_count
    stream_ids = tuple(range(font_id + 1, font_id + 1 + page_count))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] "
            f"/Count {page_count} >>"
        ).encode(),
    ]
    for stream_id in stream_ids:
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {stream_id} 0 R >>".encode()
        )
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    for text in page_texts:
        stream = b"BT /F1 9 Tf 30 750 Td (" + escaped(text) + b") Tj ET"
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    )
    return bytes(output)


def _arera_343_source(
    *, body: bytes | None = None, layout_version: str = ARERA_343_CONFIRMATION_LAYOUT_VERSION
) -> RegulatorySourceDocument:
    content = _synthetic_343_pdf() if body is None else body
    return RegulatorySourceDocument(
        acquired=AcquiredOfficialBytes(
            url=ARERA_343_PDF_URL,
            body=content,
            sha256=sha256(content).hexdigest(),
            fetched_at=datetime(2026, 10, 7, 10, tzinfo=UTC),
            content_type="application/pdf",
        ),
        source_id="arera_343",
        channel=RegulatoryRegistryChannel.ARERA_TARIFFS,
        act_id="343/2026/R/com",
        document_id="343/2026/R/com",
        published_at=date(2026, 9, 29),
        content_type="application/pdf",
        document_kind=ARERA_343_CONFIRMATION_DOCUMENT_KIND,
        layout_version=layout_version,
        record_url="https://www.arera.it/atti-e-provvedimenti/dettaglio/26/343-26",
    )


def test_arera_343_parser_emits_non_numeric_q4_confirmation_assertions() -> None:
    result = default_regulatory_parser_registry().parse(_arera_343_source())

    assert result.disposition == ParserDisposition.PARSED, result.detail
    assert result.parser_id == ARERA_343_CONFIRMATION_PARSER_ID
    assert result.parser_version == ARERA_343_CONFIRMATION_PARSER_VERSION
    assert result.facts == ()
    assert {item.component_code for item in result.effect_assertions} == {
        "ASOS",
        "ARIM",
        "UC3",
        "UC6",
    }
    for assertion in result.effect_assertions:
        assert assertion.validity == DatePeriod(start=date(2026, 10, 1), end=date(2027, 1, 1))
        assert assertion.source_sha256 == result.source_sha256
        assert assertion.raw_value_token == "sono confermati"
        assert assertion.segments == (
            AreraCustomerSegment.RESIDENT,
            AreraCustomerSegment.NON_RESIDENT,
        )


@pytest.mark.parametrize(
    ("source", "reason"),
    [
        (
            _arera_343_source(layout_version="unregistered-v1"),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _arera_343_source(body=_synthetic_343_pdf(page_count=7)),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _arera_343_source(
                body=_synthetic_343_pdf().replace(b"sono confermati", b"sono modificati")
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
    ],
)
def test_arera_343_parser_rejects_unknown_identity_or_layout(
    source: RegulatorySourceDocument, reason: RegulatoryRolloverReason
) -> None:
    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == reason
    assert result.facts == ()
    assert result.effect_assertions == ()


@pytest.mark.parametrize(
    "source",
    (
        _arera_343_source(body=_synthetic_343_pdf(cover_text="DELIBERAZIONE 342/2026/R/com")),
        _arera_343_source(
            body=_synthetic_343_pdf(final_page_text="4.2 Efficacia dal 2 ottobre 2026.")
        ),
        _arera_343_source(body=_synthetic_343_pdf(article_replace=("Articolo 1", "Articolo 2"))),
        _arera_343_source(
            body=_synthetic_343_pdf(article_replace=("1.2 I valori", "1.1 I valori"))
        ),
        _arera_343_source(body=_synthetic_343_pdf(empty_page=3)),
    ),
)
def test_arera_343_parser_rejects_incomplete_or_changed_document_sections(
    source: RegulatorySourceDocument,
) -> None:
    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
    assert result.effect_assertions == ()


def test_arera_343_parser_rejects_record_and_download_url_mismatch() -> None:
    source = _arera_343_source()
    altered = replace(source, act_id="342/2026/R/com")

    result = default_regulatory_parser_registry().parse(altered)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE


def test_arera_343_parser_returns_typed_failure_for_malformed_pdf() -> None:
    result = default_regulatory_parser_registry().parse(_arera_343_source(body=b"not a PDF"))

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.PARSER_FAILURE
    assert result.facts == ()
    assert result.effect_assertions == ()


def _adm_excise_source(
    *,
    body: bytes | None = None,
    content_type: str = "application/pdf",
    layout_version: str = ADM_EXCISE_LAYOUT_VERSION,
) -> RegulatorySourceDocument:
    source_url = (
        "https://www.adm.gov.it/portale/documents/20182/43975520/"
        "Aliquote+nazionali++aggiornamento+al+18+settembre+2026.pdf/"
        "059466a2-2c62-b667-6dd8-9474d215f88a?t=1789719130770"
    )
    content = body or _synthetic_adm_pdf(
        "Aggiornamento al 18 settembre 2026",
        "ACCISE SULL' ENERGIA ELETTRICA",
        "per qualsiasi applicazione nelle abitazioni",
        "€ 0,0227 per ogni kWh",
        "D.M. 30/12/2011",
        "Pagina 5 di 11",
    )
    return RegulatorySourceDocument(
        acquired=AcquiredOfficialBytes(
            url=source_url,
            body=content,
            sha256=sha256(content).hexdigest(),
            fetched_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
            final_url=source_url,
            content_type=content_type,
        ),
        source_id="adm_excise_20260918",
        channel=RegulatoryRegistryChannel.ADM_EXCISE,
        act_id="adm-excise:2026-09-18:059466a2-2c62-b667-6dd8-9474d215f88a",
        document_id="059466a2-2c62-b667-6dd8-9474d215f88a",
        published_at=date(2026, 9, 18),
        content_type=content_type,
        document_kind=ADM_EXCISE_DOMESTIC_KIND,
        layout_version=layout_version,
    )


def test_supported_arera_workbook_emits_source_located_decimal_facts() -> None:
    source = _source()
    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.PARSED, result.detail
    assert result.parser_id == ARERA_DOMESTIC_WORKBOOK_PARSER_ID
    assert result.parser_version == ARERA_DOMESTIC_WORKBOOK_PARSER_VERSION
    assert len(result.facts) == 36
    assert all(fact.status == VerificationStatus.VERIFIED for fact in result.facts)
    assert all(fact.source_sha256 == source.acquired.sha256 for fact in result.facts)
    expected_validity = DatePeriod(start=date(2026, 10, 1), end=date(2026, 11, 1))
    assert {fact.validity.start.month for fact in result.facts} == {10, 11, 12}
    resident_energy_network = next(
        fact
        for fact in result.facts
        if fact.segment == AreraCustomerSegment.RESIDENT
        and fact.component_code == "network_total"
        and fact.quota == BillingQuota.CONSUMPTION
        and fact.validity == expected_validity
    )
    assert resident_energy_network.value.as_tuple().exponent == -6
    assert resident_energy_network.raw_value_token == "0.01473"
    assert resident_energy_network.locator == "ottobre 2026!I14"
    assert "Spec 005" in resident_energy_network.derivation


def test_supported_adm_pdf_emits_exact_domestic_excise_fact() -> None:
    source = _adm_excise_source()

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.PARSED, result.detail
    assert result.parser_id == ADM_EXCISE_DOMESTIC_PARSER_ID
    assert result.parser_version == ADM_EXCISE_DOMESTIC_PARSER_VERSION
    assert len(result.facts) == 1
    fact = result.facts[0]
    assert fact.family == RegulatoryFactFamily.EXCISE_RATE
    assert fact.value == Decimal("0.0227")
    assert fact.unit == RateUnit.EUR_PER_KWH
    assert fact.validity == DatePeriod(start=date(2026, 9, 18), end=date.max)
    assert fact.source_id == source.source_id
    assert fact.source_sha256 == source.acquired.sha256
    assert fact.act_id == source.act_id
    assert fact.published_at == date(2026, 9, 18)
    assert fact.locator == (
        "page 5/11: ACCISE SULL' ENERGIA ELETTRICA / per qualsiasi applicazione nelle abitazioni"
    )
    assert fact.raw_value_token == "€ 0,0227 per ogni kWh"
    assert "D.M. 30/12/2011" in fact.effect.provision


@pytest.mark.parametrize(
    ("source", "reason"),
    [
        (
            _adm_excise_source(content_type="text/html"),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(layout_version="unknown-adm-pdf-v2"),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per applicazione non domestica",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 19 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.SOURCE_CHANGED_UNEXPECTEDLY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 septembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 31 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                    "€ 0,0228 per ogni kWh",
                )
            ),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                    positions_override={"per qualsiasi applicazione": (130, 474)},
                )
            ),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "Pagina 5 di 11",
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            _adm_excise_source(
                body=_synthetic_adm_pdf(
                    "Aggiornamento al 18 settembre 2026",
                    "ACCISE SULL' ENERGIA ELETTRICA",
                    "per qualsiasi applicazione nelle abitazioni",
                    "€ 0,0227 per ogni kWh",
                    "D.M. 30/12/2011",
                    "Pagina 5 di 11",
                    page_count=10,
                )
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
    ],
)
def test_adm_parser_rejects_unknown_identity_or_changed_numeric_layout(
    source: RegulatorySourceDocument,
    reason: RegulatoryRolloverReason,
) -> None:
    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == reason
    assert result.facts == ()


def test_adm_parser_accepts_a_changed_rate_only_when_exact_domestic_row_still_matches() -> None:
    source = _adm_excise_source(
        body=_synthetic_adm_pdf(
            "Aggiornamento al 18 settembre 2026",
            "ACCISE SULL' ENERGIA ELETTRICA",
            "per qualsiasi applicazione nelle abitazioni",
            "€ 0,0228 per ogni kWh",
            "D.M. 30/12/2011",
            "Pagina 5 di 11",
        )
    )

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.PARSED, result.detail
    assert result.facts[0].value == Decimal("0.0228")


def test_adm_parser_returns_typed_failure_for_malformed_pdf_bytes() -> None:
    source = _adm_excise_source(body=b"%PDF-1.4\ntruncated")

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.PARSER_FAILURE
    assert result.facts == ()


def test_adm_parser_returns_typed_failure_if_pdf_dependency_is_unavailable(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def missing_pdf_module(module_name: str) -> Any:
        raise ImportError(module_name)

    monkeypatch.setattr(
        "italian_energy.arera.rollover_parsers.importlib.import_module",
        missing_pdf_module,
    )

    result = default_regulatory_parser_registry().parse(_adm_excise_source())

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.PARSER_FAILURE
    assert result.facts == ()


def test_adm_parser_returns_typed_failure_for_pdf_without_text_coordinates(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class MalformedPage:
        def extract_text(self, *, visitor_text: Any = None) -> str:
            if visitor_text is not None:
                visitor_text("fragment", None, [1.0], None, 11.0)
            return (
                "Aggiornamento al 18 settembre 2026 ACCISE SULL' ENERGIA ELETTRICA "
                "per qualsiasi applicazione nelle abitazioni € 0,0227 per ogni kWh "
                "D.M. 30/12/2011 Pagina 5 di 11"
            )

    class MalformedReader:
        def __init__(self, _source: Any) -> None:
            self.is_encrypted = False
            self.pages = tuple(MalformedPage() for _ in range(11))

    monkeypatch.setattr("pypdf.PdfReader", MalformedReader)

    result = default_regulatory_parser_registry().parse(_adm_excise_source())

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.PARSER_FAILURE
    assert result.facts == ()


def test_adm_parser_rejects_encrypted_pdf_without_attempting_interpretation() -> None:
    writer = PdfWriter()
    for _ in range(11):
        writer.add_blank_page(width=612, height=792)
    writer.encrypt("fixture-password")
    encrypted = io.BytesIO()
    writer.write(encrypted)

    result = default_regulatory_parser_registry().parse(
        _adm_excise_source(body=encrypted.getvalue())
    )

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
    assert result.facts == ()


def test_unknown_source_layout_is_review_required_without_guessing() -> None:
    result = default_regulatory_parser_registry().parse(
        _source(layout_version="unapproved-2026-layout")
    )

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
    assert not result.facts
    assert result.parser_id is None


@pytest.mark.parametrize(
    ("table", "parser_id", "source_id", "locator", "token", "valid_from"),
    [
        (
            True,
            NORMATTIVA_VAT_TABLE_PARSER_ID,
            "vat_dpr_633",
            "bodyTesto/attachment-just-text/PARTE III/item 103",
            "103) energia elettrica per uso domestico; energia elettrica per altri usi;",
            date(2025, 12, 13),
        ),
        (
            False,
            NORMATTIVA_VAT_RATE_PARSER_ID,
            "vat_dpr_633_art16",
            "bodyTesto/art_16/art-just-text-akn/ins_3",
            "dieci per cento",
            date(2016, 1, 1),
        ),
    ],
)
def test_normattiva_vat_parsers_emit_exact_provenance_and_half_open_validity(
    table: bool,
    parser_id: str,
    source_id: str,
    locator: str,
    token: str,
    valid_from: date,
) -> None:
    source = _normattiva_vat_source(table=table)

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.PARSED, result.detail
    assert result.parser_id == parser_id
    assert len(result.facts) == 1
    fact = result.facts[0]
    assert fact.fact_id == f"{source_id}:vat-rate:{valid_from.isoformat()}:2027-01-01"
    assert fact.family.value == "vat_rate"
    assert fact.value == Decimal("10")
    assert fact.unit == RateUnit.PERCENT
    assert fact.validity == DatePeriod(start=valid_from, end=date(2027, 1, 1))
    assert fact.source_id == source_id
    assert fact.source_sha256 == source.acquired.sha256
    assert fact.source_url == source.acquired.url
    assert fact.published_at == date(1972, 10, 26)
    assert fact.fetched_at == source.acquired.fetched_at
    assert fact.parser_id == parser_id
    assert fact.parser_version == "1.0.0"
    assert fact.locator == locator
    assert fact.raw_value_token == token
    assert fact.effect.provision
    assert fact.status == VerificationStatus.VERIFIED


def test_art16_parser_requires_its_exact_supported_rate_clause() -> None:
    source = _normattiva_vat_source(table=False)
    changed = source.acquired.body.replace(b"dieci per cento", b"undici per cento")
    source = RegulatorySourceDocument(
        acquired=AcquiredOfficialBytes(
            url=source.acquired.url,
            body=changed,
            sha256=sha256(changed).hexdigest(),
            fetched_at=source.acquired.fetched_at,
            content_type=source.acquired.content_type,
        ),
        source_id=source.source_id,
        channel=source.channel,
        act_id=source.act_id,
        document_id=source.document_id,
        published_at=source.published_at,
        content_type=source.content_type,
        document_kind=source.document_kind,
        layout_version=source.layout_version,
    )

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
    assert result.facts == ()


def test_table_a_parser_requires_unique_domestic_electricity_item_103() -> None:
    source = _normattiva_vat_source(table=True)
    changed = source.acquired.body.replace(
        b"energia elettrica per uso domestico", b"energia elettrica"
    )
    source = RegulatorySourceDocument(
        acquired=AcquiredOfficialBytes(
            url=source.acquired.url,
            body=changed,
            sha256=sha256(changed).hexdigest(),
            fetched_at=source.acquired.fetched_at,
            content_type=source.acquired.content_type,
        ),
        source_id=source.source_id,
        channel=source.channel,
        act_id=source.act_id,
        document_id=source.document_id,
        published_at=source.published_at,
        content_type=source.content_type,
        document_kind=source.document_kind,
        layout_version=source.layout_version,
    )

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY
    assert result.facts == ()


def test_normattiva_vat_parser_requires_vig_date_to_match_acquisition_day() -> None:
    source = _normattiva_vat_source(table=False, url_date=date(2026, 10, 6))

    result = default_regulatory_parser_registry().parse(source)

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.SOURCE_CHANGED_UNEXPECTEDLY
    assert result.facts == ()


@pytest.mark.parametrize(
    ("body", "reason"),
    [
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(b"13-12-2025", b"13/12/2025"),
            RegulatoryRolloverReason.INVALID_VALIDITY_INTERVAL,
        ),
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(b"13-12-2025", b"31-02-2025"),
            RegulatoryRolloverReason.INVALID_VALIDITY_INTERVAL,
        ),
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(
                b'<div class="bodyTesto">', b'<div class="notBodyTesto">'
            ),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(b"13-12-2025", b"6-10-2026"),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(b"PARTE III (194)", b"PARTE II (194)"),
            RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE,
        ),
        (
            VAT_TABLE_FIXTURE.read_bytes().replace(b"104)", b"105)"),
            RegulatoryRolloverReason.AMBIGUOUS_APPLICABILITY,
        ),
        (
            b"<html><body>\xff</body></html>",
            RegulatoryRolloverReason.PARSER_FAILURE,
        ),
    ],
)
def test_table_a_parser_fails_closed_on_changed_dates_and_akn_structure(
    body: bytes,
    reason: RegulatoryRolloverReason,
) -> None:
    result = default_regulatory_parser_registry().parse(
        _normattiva_vat_source(table=True, body=body)
    )

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == reason
    assert result.facts == ()


def test_art16_parser_fails_closed_on_missing_heading_or_changed_wrapping() -> None:
    fixture = VAT_ART16_FIXTURE.read_bytes()
    changed_sources = (
        _normattiva_vat_source(
            table=False,
            body=fixture.replace(b'id="art_16"', b'id="art_15"'),
        ),
        _normattiva_vat_source(
            table=False,
            body=fixture.replace(b"((L'aliquota", b"(L'aliquota").replace(
                b"articolo 34))", b"articolo 34)"
            ),
        ),
    )

    for source in changed_sources:
        result = default_regulatory_parser_registry().parse(source)
        assert result.disposition == ParserDisposition.REVIEW_REQUIRED
        assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
        assert result.facts == ()


def test_recognized_layout_with_malformed_workbook_is_typed_parser_failure() -> None:
    result = default_regulatory_parser_registry().parse(_source(body=b"not-an-xlsx"))

    assert result.disposition == ParserDisposition.REVIEW_REQUIRED
    assert result.reason_code == RegulatoryRolloverReason.PARSER_FAILURE
    assert not result.facts


def test_parser_registry_rejects_duplicate_exact_layout_registration() -> None:
    registry = RegulatoryParserRegistry()
    parser = AreraDomesticWorkbookParser()
    registry.register(parser)

    with pytest.raises(ValueError, match="already registered"):
        registry.register(parser)


def test_wrong_media_type_is_unsupported_source() -> None:
    source = _source()
    source = RegulatorySourceDocument(
        acquired=source.acquired,
        source_id=source.source_id,
        channel=source.channel,
        act_id=source.act_id,
        document_id=source.document_id,
        published_at=source.published_at,
        content_type="application/pdf",
        document_kind=source.document_kind,
        layout_version=source.layout_version,
    )

    result = default_regulatory_parser_registry().parse(source)

    assert result.reason_code == RegulatoryRolloverReason.UNSUPPORTED_REGULATORY_SOURCE
    assert result.disposition == ParserDisposition.REVIEW_REQUIRED


def test_source_document_requires_publication_before_acquisition() -> None:
    source = _source()

    with pytest.raises(ValueError, match="before publication"):
        RegulatorySourceDocument(
            acquired=AcquiredOfficialBytes(
                url=source.acquired.url,
                body=source.acquired.body,
                sha256=source.acquired.sha256,
                fetched_at=datetime(2026, 9, 20, tzinfo=UTC),
            ),
            source_id=source.source_id,
            channel=source.channel,
            act_id=source.act_id,
            document_id=source.document_id,
            published_at=PUBLISHED_AT,
            content_type=source.content_type,
            document_kind=source.document_kind,
            layout_version=source.layout_version,
        )


def test_parser_result_rejects_incomplete_or_misattributed_facts() -> None:
    fact = default_regulatory_parser_registry().parse(_source()).facts[0]
    common = {
        "disposition": ParserDisposition.PARSED,
        "source_id": fact.source_id,
        "source_sha256": fact.source_sha256,
        "parser_id": "fixture-parser",
        "parser_version": "1.0.0",
        "facts": (fact,),
        "detail": "synthetic contract test",
    }

    with pytest.raises(ValueError, match="retain the source identity"):
        RegulatoryParseResult.model_validate({**common, "source_id": "different-source"})

    with pytest.raises(ValueError, match="cannot contain a failure reason"):
        RegulatoryParseResult.model_validate(
            {**common, "reason_code": RegulatoryRolloverReason.PARSER_FAILURE}
        )

    with pytest.raises(ValueError, match="parsed result requires parser identity"):
        RegulatoryParseResult.model_validate({**common, "parser_id": None})


def test_parser_review_result_cannot_contain_partial_facts_or_omit_reason() -> None:
    fact = default_regulatory_parser_registry().parse(_source()).facts[0]
    common = {
        "disposition": ParserDisposition.REVIEW_REQUIRED,
        "source_id": fact.source_id,
        "source_sha256": fact.source_sha256,
        "detail": "unsupported layout",
    }

    with pytest.raises(ValueError, match="requires a reason"):
        RegulatoryParseResult.model_validate(common)

    with pytest.raises(ValueError, match="cannot expose partial facts"):
        RegulatoryParseResult.model_validate(
            {
                **common,
                "reason_code": RegulatoryRolloverReason.PARSER_FAILURE,
                "facts": (fact,),
            }
        )


def test_source_document_rejects_empty_identity_and_wrong_registry_host() -> None:
    source = _source()
    with pytest.raises(ValueError, match="identity fields cannot be empty"):
        RegulatorySourceDocument(
            acquired=source.acquired,
            source_id=" ",
            channel=source.channel,
            act_id=source.act_id,
            document_id=source.document_id,
            published_at=source.published_at,
            content_type=source.content_type,
            document_kind=source.document_kind,
            layout_version=source.layout_version,
        )

    wrong_registry_source = AcquiredOfficialBytes(
        url="https://www.adm.gov.it/fixture.xlsx",
        body=source.acquired.body,
        sha256=source.acquired.sha256,
        fetched_at=source.acquired.fetched_at,
    )
    with pytest.raises(ValueError, match="match its official registry channel"):
        RegulatorySourceDocument(
            acquired=wrong_registry_source,
            source_id=source.source_id,
            channel=source.channel,
            act_id=source.act_id,
            document_id=source.document_id,
            published_at=source.published_at,
            content_type=source.content_type,
            document_kind=source.document_kind,
            layout_version=source.layout_version,
        )
